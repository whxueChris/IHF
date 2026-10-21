from __future__ import annotations

import math
from pathlib import Path
from statistics import mean

from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont


# ---------------------------------------------------------------------------
# Files and export geometry
# ---------------------------------------------------------------------------

SCRIPT_DIR = Path(__file__).resolve().parent
SOURCE_DATA = SCRIPT_DIR / "figure6c_localization_matched_vs_mismatch_source_data.xlsx"
OUTPUT_PNG = SCRIPT_DIR / "figure6c_localization_matched_vs_mismatch_transparent.png"

EXPORT_DPI = 600
FIGURE_WIDTH_MM = 54.0
FIGURE_HEIGHT_MM = 62.0
Y_MIN = 0.0
Y_MAX = 25.0


# ---------------------------------------------------------------------------
# Publication-oriented visual style
# ---------------------------------------------------------------------------

INK = "#202020"
AXIS = "#585858"
MATCHED = "#1769AA"
MATCHED_LIGHT = "#A9CAE1"
MISMATCHED = "#C9796C"
MISMATCHED_LIGHT = "#F1B8AC"
PAIR_LINE = "#B9B9B9"


def mm_to_px(value_mm: float) -> int:
    return round(value_mm / 25.4 * EXPORT_DPI)


def pt_to_px(value_pt: float) -> int:
    return max(1, round(value_pt / 72.0 * EXPORT_DPI))


def load_font(size_pt: float, bold: bool = False) -> ImageFont.FreeTypeFont:
    candidates = [
        Path(r"C:\Windows\Fonts\arialbd.ttf")
        if bold
        else Path(r"C:\Windows\Fonts\arial.ttf"),
        Path(r"C:\Windows\Fonts\calibrib.ttf")
        if bold
        else Path(r"C:\Windows\Fonts\calibri.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size=pt_to_px(size_pt))
    return ImageFont.load_default(size=pt_to_px(size_pt))


FONT_AXIS_LABEL = load_font(6.8)
FONT_TICK = load_font(6.5)
FONT_CATEGORY = load_font(6.8, True)
FONT_MEAN = load_font(6.8, True)
FONT_P = load_font(6.5)



# ---------------------------------------------------------------------------
# Data and exact paired test
# ---------------------------------------------------------------------------

def read_participant_means() -> list[dict[str, float | str | int]]:
    workbook = load_workbook(SOURCE_DATA, data_only=True, read_only=True)
    sheet = workbook["Participant Means"]
    rows = list(sheet.iter_rows(values_only=True))
    headers = [str(value) for value in rows[0]]
    records = [dict(zip(headers, row)) for row in rows[1:] if row[0]]
    workbook.close()

    expected_ids = ["P002", "P003", "P005", "P006", "P007", "P012", "P014"]
    observed_ids = [str(record["participant_id"]) for record in records]
    if observed_ids != expected_ids:
        raise ValueError(f"Unexpected participant order: {observed_ids}")
    if len(records) != 7:
        raise ValueError(f"Expected 7 participants, found {len(records)}")
    if any(int(record["matched_n"]) != 9 for record in records):
        raise ValueError("Every participant must have 9 matched targets")
    if any(int(record["mismatch_n"]) != 9 for record in records):
        raise ValueError("Every participant must have 9 mismatch targets")
    return records


def rank_absolute_differences(differences: list[float]) -> list[float]:
    indexed = sorted(enumerate(abs(value) for value in differences), key=lambda item: item[1])
    ranks = [0.0] * len(differences)
    cursor = 0
    while cursor < len(indexed):
        end = cursor + 1
        while end < len(indexed) and math.isclose(
            indexed[end][1], indexed[cursor][1], rel_tol=0.0, abs_tol=1e-12
        ):
            end += 1
        average_rank = ((cursor + 1) + end) / 2.0
        for position in range(cursor, end):
            ranks[indexed[position][0]] = average_rank
        cursor = end
    return ranks


def exact_two_sided_wilcoxon(differences: list[float]) -> float:
    nonzero = [value for value in differences if not math.isclose(value, 0.0, abs_tol=1e-12)]
    ranks = rank_absolute_differences(nonzero)
    observed_w_plus = sum(
        rank for rank, difference in zip(ranks, nonzero) if difference > 0
    )
    total_rank = sum(ranks)
    observed_deviation = abs(observed_w_plus - total_rank / 2.0)

    permutations = 2 ** len(ranks)
    at_least_as_extreme = 0
    for mask in range(permutations):
        w_plus = sum(rank for index, rank in enumerate(ranks) if (mask >> index) & 1)
        if abs(w_plus - total_rank / 2.0) >= observed_deviation - 1e-12:
            at_least_as_extreme += 1
    return at_least_as_extreme / permutations


# ---------------------------------------------------------------------------
# Drawing helpers
# ---------------------------------------------------------------------------

def text_size(
    draw: ImageDraw.ImageDraw, text: str, font: ImageFont.ImageFont
) -> tuple[int, int]:
    box = draw.textbbox((0, 0), text, font=font)
    return box[2] - box[0], box[3] - box[1]


def draw_centered(
    draw: ImageDraw.ImageDraw,
    x: float,
    y: float,
    text: str,
    font: ImageFont.ImageFont,
    fill: str,
) -> None:
    width, _ = text_size(draw, text, font)
    draw.text((x - width / 2.0, y), text, font=font, fill=fill)


def draw_right(
    draw: ImageDraw.ImageDraw,
    x: float,
    y: float,
    text: str,
    font: ImageFont.ImageFont,
    fill: str,
) -> None:
    width, _ = text_size(draw, text, font)
    draw.text((x - width, y), text, font=font, fill=fill)


def draw_circle(
    draw: ImageDraw.ImageDraw,
    center: tuple[float, float],
    radius: float,
    fill: str,
    outline: str,
    width: int,
) -> None:
    x, y = center
    draw.ellipse(
        (x - radius, y - radius, x + radius, y + radius),
        fill=fill,
        outline=outline,
        width=width,
    )


def draw_rotated_text(
    image: Image.Image,
    center: tuple[float, float],
    text: str,
    font: ImageFont.ImageFont,
    fill: str,
) -> None:
    temporary_draw = ImageDraw.Draw(image)
    width, height = text_size(temporary_draw, text, font)
    label = Image.new("RGBA", (width + pt_to_px(2), height + pt_to_px(2)), (0, 0, 0, 0))
    label_draw = ImageDraw.Draw(label)
    label_draw.text((pt_to_px(1), pt_to_px(1)), text, font=font, fill=fill)
    rotated = label.rotate(90, expand=True, resample=Image.Resampling.BICUBIC)
    image.alpha_composite(
        rotated,
        (
            round(center[0] - rotated.width / 2),
            round(center[1] - rotated.height / 2),
        ),
    )


# ---------------------------------------------------------------------------
# Figure
# ---------------------------------------------------------------------------

def render_figure(records: list[dict[str, float | str | int]]) -> Image.Image:
    width = mm_to_px(FIGURE_WIDTH_MM)
    height = mm_to_px(FIGURE_HEIGHT_MM)
    image = Image.new("RGBA", (width, height), (255, 255, 255, 0))
    draw = ImageDraw.Draw(image)

    plot_left = mm_to_px(15.3)
    plot_right = width - mm_to_px(2.0)
    plot_top = mm_to_px(5.8)
    plot_bottom = height - mm_to_px(10.5)

    def y_map(value: float) -> float:
        return plot_bottom - (value - Y_MIN) / (Y_MAX - Y_MIN) * (
            plot_bottom - plot_top
        )

    axis_width = pt_to_px(0.55)
    tick_width = pt_to_px(0.45)

    draw.line(
        (plot_left, plot_top, plot_left, plot_bottom),
        fill=AXIS,
        width=axis_width,
    )
    draw.line(
        (plot_left, plot_bottom, plot_right, plot_bottom),
        fill=AXIS,
        width=axis_width,
    )

    for tick in (0, 5, 10, 15, 20, 25):
        y = y_map(float(tick))
        draw.line(
            (plot_left - pt_to_px(2.0), y, plot_left, y),
            fill=AXIS,
            width=tick_width,
        )
        _, tick_height = text_size(draw, str(tick), FONT_TICK)
        draw_right(
            draw,
            plot_left - pt_to_px(3.1),
            y - tick_height / 2.0,
            str(tick),
            FONT_TICK,
            AXIS,
        )

    draw_rotated_text(
        image,
        (mm_to_px(4.8), (plot_top + plot_bottom) / 2.0),
        "Mean localization error (%)",
        FONT_AXIS_LABEL,
        INK,
    )

    axis_span = plot_right - plot_left
    matched_x = plot_left + axis_span * 0.34
    # Keep the mismatch group sufficiently far from the right canvas edge so
    # its direct mean label can sit beyond the mean bar without overlapping it.
    mismatch_x = plot_left + axis_span * 0.70
    participant_jitter = [
        mm_to_px(value) for value in (-1.5, -1.0, -0.5, 0.0, 0.5, 1.0, 1.5)
    ]
    dot_radius = pt_to_px(1.8)

    matched_values = [float(record["matched_mean_error_pct"]) for record in records]
    mismatch_values = [float(record["mismatch_mean_error_pct"]) for record in records]
    differences = [
        mismatch - matched for matched, mismatch in zip(matched_values, mismatch_values)
    ]

    for matched, mismatch, jitter in zip(
        matched_values, mismatch_values, participant_jitter
    ):
        draw.line(
            (
                matched_x + jitter,
                y_map(matched),
                mismatch_x + jitter,
                y_map(mismatch),
            ),
            fill=PAIR_LINE,
            width=pt_to_px(0.55),
        )
        draw_circle(
            draw,
            (matched_x + jitter, y_map(matched)),
            dot_radius,
            MATCHED_LIGHT,
            MATCHED,
            pt_to_px(0.45),
        )
        draw_circle(
            draw,
            (mismatch_x + jitter, y_map(mismatch)),
            dot_radius,
            MISMATCHED_LIGHT,
            MISMATCHED,
            pt_to_px(0.45),
        )

    matched_mean = mean(matched_values)
    mismatch_mean = mean(mismatch_values)
    if round(matched_mean, 1) != 5.9 or round(mismatch_mean, 1) != 12.2:
        raise ValueError(
            f"Unexpected means: matched={matched_mean:.6f}, mismatch={mismatch_mean:.6f}"
        )

    mean_half_width = mm_to_px(3.0)
    mean_line_width = pt_to_px(1.0)
    matched_mean_y = y_map(matched_mean)
    mismatch_mean_y = y_map(mismatch_mean)
    draw.line(
        (
            matched_x - mean_half_width,
            matched_mean_y,
            matched_x + mean_half_width,
            matched_mean_y,
        ),
        fill=MATCHED,
        width=mean_line_width,
    )
    draw.line(
        (
            mismatch_x - mean_half_width,
            mismatch_mean_y,
            mismatch_x + mean_half_width,
            mismatch_mean_y,
        ),
        fill=MISMATCHED,
        width=mean_line_width,
    )

    _, mean_label_height = text_size(draw, "5.9%", FONT_MEAN)
    draw_right(
        draw,
        matched_x - mean_half_width - mm_to_px(1.0),
        matched_mean_y - mean_label_height / 2.0,
        f"{matched_mean:.1f}%",
        FONT_MEAN,
        MATCHED,
    )

    mismatch_label = f"{mismatch_mean:.1f}%"
    label_width, label_height = text_size(draw, mismatch_label, FONT_MEAN)

    preferred_label_x = mismatch_x + mean_half_width + mm_to_px(1.0)
    safe_label_x = min(
        preferred_label_x,
        width - label_width - mm_to_px(1.0),
    )

    draw.text(
        (
            safe_label_x,
            mismatch_mean_y - label_height / 2.0,
        ),
        mismatch_label,
        font=FONT_MEAN,
        fill=MISMATCHED,
    )

    p_value = exact_two_sided_wilcoxon(differences)
    if not math.isclose(p_value, 0.015625, rel_tol=0.0, abs_tol=1e-12):
        raise ValueError(f"Unexpected exact paired P value: {p_value}")

    highest_value = max(matched_values + mismatch_values)
    bracket_value = min(Y_MAX - 3.0, highest_value + 2.0)
    p_text_value = bracket_value + 1.5
    bracket_y = y_map(bracket_value)
    bracket_drop = mm_to_px(1.0)
    draw.line(
        (matched_x, bracket_y, mismatch_x, bracket_y),
        fill=INK,
        width=axis_width,
    )
    draw.line(
        (matched_x, bracket_y, matched_x, bracket_y + bracket_drop),
        fill=INK,
        width=axis_width,
    )
    draw.line(
        (mismatch_x, bracket_y, mismatch_x, bracket_y + bracket_drop),
        fill=INK,
        width=axis_width,
    )
    _, p_height = text_size(draw, "P = 0.0156", FONT_P)
    draw_centered(
        draw,
        (matched_x + mismatch_x) / 2.0,
        y_map(p_text_value) - p_height / 2.0,
        "P = 0.0156",
        FONT_P,
        INK,
    )

    label_y = plot_bottom + mm_to_px(2.3)
    draw_centered(draw, matched_x, label_y, "Matched", FONT_CATEGORY, MATCHED)
    draw_centered(
        draw, mismatch_x, label_y, "Mismatch", FONT_CATEGORY, MISMATCHED
    )

    return image


def main() -> None:
    records = read_participant_means()
    figure = render_figure(records)
    figure.save(
        OUTPUT_PNG,
        format="PNG",
        dpi=(EXPORT_DPI, EXPORT_DPI),
        optimize=True,
    )
    print(f"Saved: {OUTPUT_PNG}")


if __name__ == "__main__":
    main()
